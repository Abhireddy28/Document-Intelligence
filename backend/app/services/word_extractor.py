import docx
import logging
from typing import Dict, Any, List

logger = logging.getLogger("agent64.word_extractor")

class WordExtractor:
    """Extracts text, headings, paragraphs, and tables from Word (.docx) documents."""

    def extract(self, file_path: str) -> Dict[str, Any]:
        doc = docx.Document(file_path)
        
        paragraphs = []
        headings = []
        raw_text_lines = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style and "Heading" in p.style.name:
                headings.append({"style": p.style.name, "text": text})
            else:
                paragraphs.append(text)
            raw_text_lines.append(text)

        tables_data = []
        for t_idx, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(cells)
            if rows:
                headers = rows[0]
                table_rows = []
                for r in rows[1:]:
                    row_dict = {headers[i] if i < len(headers) else f"col_{i}": cell for i, cell in enumerate(r)}
                    table_rows.append(row_dict)
                tables_data.append({
                    "table_index": t_idx + 1,
                    "headers": headers,
                    "rows": table_rows
                })

        return {
            "headings": headings,
            "paragraphs": paragraphs,
            "tables": tables_data,
            "raw_text": "\n\n".join(raw_text_lines),
            "extraction_method": "DOCX_TEXT"
        }

word_extractor = WordExtractor()
