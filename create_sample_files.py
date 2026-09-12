import os
import shutil
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from docx import Document
from PIL import Image, ImageDraw, ImageFont

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SAMPLE_DIR = os.path.join(BASE_DIR, "sample_documents")
UPLOADS_DIR = os.path.join(BASE_DIR, "backend", "uploads")

os.makedirs(os.path.join(SAMPLE_DIR, "marks_card"), exist_ok=True)
os.makedirs(os.path.join(SAMPLE_DIR, "attendance"), exist_ok=True)
os.makedirs(os.path.join(SAMPLE_DIR, "certificates"), exist_ok=True)
os.makedirs(os.path.join(SAMPLE_DIR, "circulars"), exist_ok=True)
os.makedirs(UPLOADS_DIR, exist_ok=True)

def generate_clean_marks_pdf(file_path):
    c = canvas.Canvas(file_path, pagesize=letter)
    width, height = letter
    
    # Header banner
    c.setFillColor(colors.HexColor("#0B1730"))
    c.rect(0, height - 70, width, 70, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(width / 2.0, height - 35, "VIGNAN'S UNIVERSITY")
    c.setFont("Helvetica", 11)
    c.drawCentredString(width / 2.0, height - 55, "OFFICE OF THE CONTROLLER OF EXAMINATIONS - GRADE MEMORANDUM")

    # Student details card
    c.setFillColor(colors.HexColor("#0B1730"))
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, height - 105, "Student Information:")
    
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 125, "Student Name: Rahul Kumar")
    c.drawString(320, height - 125, "Roll Number: 22CS101")
    c.drawString(50, height - 145, "Program: B.Tech Computer Science & Engineering")
    c.drawString(320, height - 145, "Semester: VI (Sixth Semester)")
    c.drawString(50, height - 165, "Academic Year: 2025-2026")
    c.drawString(320, height - 165, "Month & Year of Exam: March 2026")

    # Table header
    y = height - 200
    c.setFillColor(colors.HexColor("#426FA8"))
    c.rect(50, y, width - 100, 22, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(60, y + 6, "Code")
    c.drawString(130, y + 6, "Subject Name")
    c.drawString(320, y + 6, "Internal")
    c.drawString(380, y + 6, "External")
    c.drawString(440, y + 6, "Total")
    c.drawString(500, y + 6, "Grade")

    # Table rows
    subjects = [
        ("CS301", "Cloud Computing & DevOps", "28", "64", "92", "A+"),
        ("CS302", "Artificial Intelligence & Agents", "27", "63", "90", "A+"),
        ("CS303", "Compiler Design & Automata", "26", "60", "86", "A"),
        ("CS304", "Computer Networks & Security", "25", "58", "83", "A"),
    ]

    c.setFillColor(colors.HexColor("#0B1730"))
    c.setFont("Helvetica", 9)
    for row in subjects:
        y -= 22
        c.setStrokeColor(colors.HexColor("#D9E2EF"))
        c.line(50, y, width - 50, y)
        c.drawString(60, y + 6, row[0])
        c.drawString(130, y + 6, row[1])
        c.drawString(330, y + 6, row[2])
        c.drawString(390, y + 6, row[3])
        c.drawString(450, y + 6, row[4])
        c.drawString(505, y + 6, row[5])

    # Summary box
    y -= 40
    c.setFillColor(colors.HexColor("#E4EFFC"))
    c.rect(50, y, width - 100, 30, fill=1)
    c.setFillColor(colors.HexColor("#0B1730"))
    c.setFont("Helvetica-Bold", 10)
    c.drawString(60, y + 10, "Total Marks: 351 / 400")
    c.drawString(240, y + 10, "Percentage: 87.75%")
    c.drawString(420, y + 10, "Result: FIRST CLASS WITH DISTINCTION")

    # Signatures
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(60, 60, "Verified by: Section Officer")
    c.drawString(width - 220, 60, "Controller of Examinations")

    c.showPage()
    c.save()

def generate_noisy_marks_pdf(file_path):
    # Generates scanned-style marks card with OCR noise (22CS10I)
    c = canvas.Canvas(file_path, pagesize=letter)
    width, height = letter

    c.setFillColor(colors.HexColor("#1A202C"))
    c.setFont("Helvetica-Bold", 14)
    c.drawCentredString(width / 2.0, height - 50, "VIGNAN'S UNIVERSITY - GRADE REPORT")
    
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 90, "Student Name: Rahul Kumar")
    c.drawString(320, height - 90, "Roll Number: 22CS10I")
    c.drawString(50, height - 110, "Semester: VI")
    c.drawString(320, height - 110, "Academic Year: 2025-2026")

    y = height - 150
    c.setFillColor(colors.HexColor("#426FA8"))
    c.rect(50, y, width - 100, 20, fill=1)
    c.setFillColor(colors.white)
    c.drawString(60, y + 5, "Code")
    c.drawString(130, y + 5, "Subject Name")
    c.drawString(320, y + 5, "Internal")
    c.drawString(380, y + 5, "External")
    c.drawString(440, y + 5, "Total")
    c.drawString(500, y + 5, "Grade")

    c.setFillColor(colors.black)
    y -= 20
    c.drawString(60, y + 5, "CS301")
    c.drawString(130, y + 5, "Cloud Computing")
    c.drawString(330, y + 5, "28")
    c.drawString(390, y + 5, "64")
    c.drawString(450, y + 5, "92")
    c.drawString(505, y + 5, "A+")

    c.showPage()
    c.save()

def generate_attendance_excel(file_path):
    data = {
        "Roll Number": ["22CS101", "22CS102", "22CS103", "22CS104", "22CS105", "22IT101", "22IT102", "22ECE101", "22ECE102", "22AI101"],
        "Student Name": ["Rahul Kumar", "Priya Sharma", "Ananya Reddy", "Vikramaditya Rao", "Sneha Patel", "Karthik Varma", "Deepa Nair", "Rohan Gupta", "Meera Joshi", "Arjun Krishna"],
        "Semester": [6, 6, 6, 6, 6, 6, 6, 6, 6, 6],
        "Total Classes": [120, 120, 120, 120, 120, 120, 120, 120, 120, 120],
        "Present": [108, 114, 102, 98, 116, 104, 110, 106, 112, 115],
        "Absent": [12, 6, 18, 22, 4, 16, 10, 14, 8, 5],
        "Attendance %": [90.0, 95.0, 85.0, 81.67, 96.67, 86.67, 91.67, 88.33, 93.33, 95.83]
    }
    df = pd.DataFrame(data)
    df.to_excel(file_path, index=False)
    # Also CSV
    csv_path = file_path.replace(".xlsx", ".csv")
    df.to_csv(csv_path, index=False)

def generate_certificate_pdf(file_path):
    c = canvas.Canvas(file_path, pagesize=letter)
    width, height = letter
    
    # Border
    c.setStrokeColor(colors.HexColor("#426FA8"))
    c.setLineWidth(3)
    c.rect(30, 30, width - 60, height - 60)
    c.rect(35, 35, width - 70, height - 70)

    c.setFillColor(colors.HexColor("#0B1730"))
    c.setFont("Helvetica-Bold", 20)
    c.drawCentredString(width / 2.0, height - 100, "VIGNAN'S UNIVERSITY")
    
    c.setFont("Helvetica", 14)
    c.drawCentredString(width / 2.0, height - 130, "CERTIFICATE OF ACADEMIC EXCELLENCE")

    c.setFont("Helvetica-Oblique", 12)
    c.drawCentredString(width / 2.0, height - 180, "This is to certify that")

    c.setFont("Helvetica-Bold", 16)
    c.setFillColor(colors.HexColor("#5A3A8B"))
    c.drawCentredString(width / 2.0, height - 215, "RAHUL KUMAR (Roll No: 22CS101)")

    c.setFillColor(colors.HexColor("#0B1730"))
    c.setFont("Helvetica", 11)
    c.drawCentredString(width / 2.0, height - 260, "has been awarded First Prize for outstanding performance in")
    c.drawCentredString(width / 2.0, height - 280, "B.Tech Computer Science & Engineering - Semester VI Examination.")

    c.setFont("Helvetica", 9)
    c.drawString(60, 90, "Certificate No: VUG/2026/CS/0842")
    c.drawString(60, 75, "Date of Issue: March 15, 2026")
    c.drawRightString(width - 60, 75, "Dean, Academic Affairs")

    c.showPage()
    c.save()

def generate_circular_docx(file_path):
    doc = Document()
    doc.add_heading("VIGNAN'S UNIVERSITY", level=0)
    doc.add_heading("OFFICE OF THE CONTROLLER OF EXAMINATIONS", level=2)
    
    doc.add_paragraph("Ref No: VU/COE/2026/CIR-042\nDate: April 10, 2026")
    doc.add_heading("CIRCULAR: End Semester Examination Schedule - Spring 2026", level=1)
    
    doc.add_paragraph(
        "All Head of Departments and eligible students of B.Tech VI Semester are hereby informed "
        "that the End Semester Theory and Practical Examinations for Spring 2026 will commence from May 02, 2026."
    )
    
    t = doc.add_table(rows=1, cols=2)
    hdr_cells = t.rows[0].cells
    hdr_cells[0].text = 'Event'
    hdr_cells[1].text = 'Scheduled Date'
    
    events = [
        ("Hall Ticket Generation", "April 25, 2026"),
        ("Practical Examinations", "May 02 - May 08, 2026"),
        ("Theory Examinations", "May 10 - May 24, 2026"),
        ("Results Announcement", "June 10, 2026")
    ]
    for ev, dt in events:
        row = t.add_row().cells
        row[0].text = ev
        row[1].text = dt
        
    doc.save(file_path)

if __name__ == "__main__":
    clean_marks_path = os.path.join(SAMPLE_DIR, "marks_card", "clean_marks_card_22CS101.pdf")
    noisy_marks_path = os.path.join(SAMPLE_DIR, "marks_card", "scanned_marks_card_noisy_22CS10I.pdf")
    att_path = os.path.join(SAMPLE_DIR, "attendance", "attendance_sem6_cse.xlsx")
    cert_path = os.path.join(SAMPLE_DIR, "certificates", "merit_award_certificate.pdf")
    circ_path = os.path.join(SAMPLE_DIR, "circulars", "exam_schedule_circular_2026.docx")

    generate_clean_marks_pdf(clean_marks_path)
    generate_noisy_marks_pdf(noisy_marks_path)
    generate_attendance_excel(att_path)
    generate_certificate_pdf(cert_path)
    generate_circular_docx(circ_path)

    # Copy to uploads
    shutil.copy(clean_marks_path, os.path.join(UPLOADS_DIR, "clean_marks_card_22CS101.pdf"))
    shutil.copy(noisy_marks_path, os.path.join(UPLOADS_DIR, "scanned_marks_card_noisy_22CS10I.pdf"))
    shutil.copy(att_path, os.path.join(UPLOADS_DIR, "attendance_sem6_cse.xlsx"))
    shutil.copy(cert_path, os.path.join(UPLOADS_DIR, "merit_award_certificate.pdf"))
    shutil.copy(circ_path, os.path.join(UPLOADS_DIR, "exam_schedule_circular_2026.docx"))

    print("Sample test fixtures successfully generated and synced to uploads!")
